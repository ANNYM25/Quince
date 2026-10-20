#!/usr/bin/env python3
"""Self-contained HTML viewer for a PPTX-path deck (the HTML path makes its own with render.js).

Usage:  python3 viewer.py deck.pdf deck.pptx out.html [--title "Deck name"]

One file, slides embedded as JPG (1600 px), arrow keys / click to move, #n deep links,
speaker notes toggled with N, print-friendly. Titles come from the 'ad-title' shapes.
"""
import sys, os, base64, glob, json, subprocess, tempfile, html
from pptx import Presentation


def main(a):
    pdf, pptx, out = a[0], a[1], a[2]
    title = a[a.index('--title') + 1] if '--title' in a else os.path.splitext(os.path.basename(pptx))[0]
    tmp = tempfile.mkdtemp()
    subprocess.run(['pdftoppm', '-jpeg', '-jpegopt', 'quality=88', '-scale-to-x', '1600', '-scale-to-y', '900', pdf, os.path.join(tmp, 's')], check=True)
    imgs = ['data:image/jpeg;base64,' + base64.b64encode(open(f, 'rb').read()).decode() for f in sorted(glob.glob(os.path.join(tmp, 's-*.jpg')))]
    prs = Presentation(pptx)
    meta = []
    for i, s in enumerate(prs.slides, 1):
        t = next((sh.text_frame.text.strip() for sh in s.shapes if sh.name == 'ad-title'), '') or (title if i == 1 else f'Slide {i}')
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
        meta.append({'t': t, 'n': notes})
    doc = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title>
<style>
:root{{--navy:#0A1F44;--blue:#1D4ED8;--ice:#EFF6FF}}
*{{box-sizing:border-box;margin:0}}
body{{background:#0B1220;color:#E2E8F0;font:14px Inter,system-ui,sans-serif;height:100vh;display:flex;flex-direction:column;overflow:hidden}}
#stage{{flex:1;display:flex;align-items:center;justify-content:center;padding:16px}}
#stage img{{max-width:100%;max-height:100%;border-radius:12px;box-shadow:0 20px 60px rgba(0,0,0,.45)}}
#bar{{display:flex;gap:12px;align-items:center;justify-content:center;padding:10px 16px}}
#bar button{{background:#1E293B;color:#E2E8F0;border:0;border-radius:999px;padding:6px 14px;font:inherit;cursor:pointer}}
#bar span{{min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;max-width:60vw}}
#notes{{display:none;background:#111827;border-radius:12px;margin:0 16px 12px;padding:12px 16px;max-height:22vh;overflow:auto;line-height:1.5}}
body.n #notes{{display:block}}
@media print{{body{{display:block;height:auto;background:#fff}}#bar,#notes{{display:none!important}}#stage{{display:block;padding:0}}#stage img{{display:none}}.p{{display:block!important;width:100%;page-break-after:always;border-radius:0;box-shadow:none}}}}
</style></head><body>
<div id="stage"><img id="im" alt=""></div>
<div id="notes"></div>
<div id="bar"><button id="pv" aria-label="Previous">←</button><span id="lab"></span><button id="nx" aria-label="Next">→</button><button id="nb">Notes (N)</button></div>
<div style="display:none" id="print"></div>
<script>
const S={json.dumps(imgs)};const M={json.dumps(meta, ensure_ascii=False)};
let i=Math.max(0,Math.min(S.length-1,(parseInt(location.hash.slice(1))||1)-1));
const im=document.getElementById('im'),lab=document.getElementById('lab'),no=document.getElementById('notes');
function show(){{i=Math.max(0,Math.min(S.length-1,i));im.src=S[i];im.alt=M[i].t;lab.textContent=(i+1)+' / '+S.length+'  ·  '+M[i].t;no.textContent=M[i].n||'No speaker notes.';history.replaceState(null,'','#'+(i+1));}}
document.getElementById('pv').onclick=()=>{{i--;show()}};document.getElementById('nx').onclick=()=>{{i++;show()}};
document.getElementById('nb').onclick=()=>document.body.classList.toggle('n');
addEventListener('keydown',e=>{{if(['ArrowRight','PageDown',' '].includes(e.key)){{i++;show()}}if(['ArrowLeft','PageUp'].includes(e.key)){{i--;show()}}if(e.key==='n'||e.key==='N')document.body.classList.toggle('n')}});
im.onclick=e=>{{i+=e.offsetX>im.clientWidth/2?1:-1;show()}};
const pr=document.getElementById('print');S.forEach(s=>{{const x=document.createElement('img');x.src=s;x.className='p';pr.appendChild(x)}});
addEventListener('beforeprint',()=>pr.style.display='block');addEventListener('afterprint',()=>pr.style.display='none');
show();
</script></body></html>"""
    open(out, 'w').write(doc)
    print(f'[viewer] {len(imgs)} slides -> {out} ({len(doc) // 1024} KB)')


if __name__ == '__main__':
    if len(sys.argv) < 4: print(__doc__); sys.exit(2)
    main(sys.argv[1:])
